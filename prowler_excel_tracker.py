#!/usr/bin/env python3
"""
prowler_excel_tracker.py
========================

Build a team **assignment tracker** (Excel) from Prowler findings.

This is the spreadsheet companion to ``prowler_report.py``. It reads the same
Prowler output, groups rows into findings (one per check ID + check title) and
writes a styled ``.xlsx`` tracker you can hand to your team.

Columns produced:

    No. | Check ID | Issue Title | Severity | Assign | Status | Comments

* **Severity** is added automatically (Critical / High / Medium / Low / Informational).
* **Assign** is a dropdown populated with your team members.
* **Status** is a dropdown (Pending / Validated / False Positive / Escalated).
* **Comments** is left blank for the reviewer.
* Column **No.** uses an auto-numbering formula.

Add people to the Assign dropdown with flags
--------------------------------------------
    --assignee KT --assignee Aqmal          # repeat the flag as many times as you like
    --assignees "KT,Aqmal,Shalitha"          # or pass a comma separated list

You can also pre-fill the Assign column:

    --assign-all Aqmal                       # give every finding to one person
    --round-robin                            # spread findings across the team

Requires ``openpyxl`` (``pip install openpyxl``). The report generator
(``prowler_report.py``) must sit in the same folder.

Examples
--------
    python prowler_excel_tracker.py
    python prowler_excel_tracker.py output/prowler-output-XXXX.csv -o tracker.xlsx
    python prowler_excel_tracker.py output --assignee KT --assignee Aqmal --round-robin
    python prowler_excel_tracker.py output --assignees "KT,Aqmal,Shalitha,Keshalya,Laviru" \
        --status FAIL --open

Run ``python prowler_excel_tracker.py --help`` for all options.
"""

from __future__ import annotations

import argparse
import sys
import webbrowser
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError:  # pragma: no cover
    sys.exit(
        "[!] This script needs 'openpyxl'.\n"
        "    Install it with:  pip install openpyxl"
    )

# The report generator lives next to this file; reuse its parsers.
try:
    import prowler_report as pr
except ImportError:  # pragma: no cover
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import prowler_report as pr


# ---------------------------------------------------------------------------
# Defaults
# ---------------------------------------------------------------------------

DEFAULT_TEAM = ["KT", "Aqmal", "Shalitha", "Keshalya", "Laviru"]
STATUSES = ["Pending", "Validated", "False Positive", "Escalated"]
SEVERITIES = ["Critical", "High", "Medium", "Low", "Informational"]
DEFAULT_STATUS = "Pending"

SEVERITY_LABEL = {
    "critical": "Critical",
    "high": "High",
    "medium": "Medium",
    "low": "Low",
    "informational": "Informational",
}

HEADERS = ["No.", "Check ID", "Issue Title", "Severity", "Assign", "Status", "Comments"]
COLS = list("BCDEFGH")  # A is a thin spacer column

# --- styling (matches the "Prowler Issues Tracker" template) --------------
FONT_TITLE = Font(name="Roboto", size=16, bold=True, color="FF1A1B20")
FONT_SUB = Font(name="Roboto", size=10, italic=True, color="FF5A5F71")
FONT_LABEL = Font(name="Roboto", size=9, bold=True, color="FF414658")
FONT_VALUE = Font(name="Roboto", size=16, bold=True, color="FF1E3989")
FONT_HEAD = Font(name="Roboto", size=10, bold=True, color="FFFFFFFF")
FONT_DATA = Font(name="Roboto", size=10, color="FF1A1B20")
FONT_DATA_B = Font(name="Roboto", size=10, bold=True, color="FF1A1B20")

FILL_LABEL = PatternFill("solid", fgColor="FFEDECF4")
FILL_VALUE = PatternFill("solid", fgColor="FFF9F8FF")
FILL_HEAD = PatternFill("solid", fgColor="FF2B3A6B")
FILL_BAND = PatternFill("solid", fgColor="FFF8F9FB")
FILL_NONE = PatternFill(fill_type=None)

BORDER = Border(
    left=Side(style="thin", color="FFD1D4DA"),
    right=Side(style="thin", color="FFE1E8EF"),
    top=Side(style="thin", color="FFD1D4DA"),
    bottom=Side(style="thin", color="FFE1E8EF"),
)

AL_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

COL_WIDTHS = {"A": 3.63, "B": 10.75, "C": 22.63, "D": 42.63,
              "E": 12.63, "F": 15.13, "G": 17.63, "H": 52.63}

TITLE_TEXT = "Cloud Security Finding & Issue Tracker"
SUBTITLE_TEXT = "Continuous compliance monitoring & remediation tracking"


# ---------------------------------------------------------------------------
# Workbook builder
# ---------------------------------------------------------------------------

def build_workbook(findings, team, assign_map=None, status=DEFAULT_STATUS,
                   sheet_name="Sheet1", title=TITLE_TEXT, subtitle=SUBTITLE_TEXT):
    """Create the tracker workbook and return (workbook, worksheet)."""
    assign_map = assign_map or {}
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    first_row = 9
    last_row = first_row + len(findings) - 1 if findings else first_row

    # --- title area ---
    ws["B2"] = title
    ws["B2"].font = FONT_TITLE
    ws["B3"] = subtitle
    ws["B3"].font = FONT_SUB

    labels = ["Total Findings", "Pending", "Validated", "Escalated", "False Positive"]
    for i, text in enumerate(labels):
        c = ws.cell(row=4, column=2 + i)
        c.value = text
        c.font = FONT_LABEL
        c.fill = FILL_LABEL
        c.border = BORDER
        c.alignment = AL_CENTER

    status_col = "G"
    formulas = [
        f"=COUNT(B{first_row}:B{last_row})",
        f'=COUNTIF({status_col}{first_row}:{status_col}{last_row}, "Pending")',
        f'=COUNTIF({status_col}{first_row}:{status_col}{last_row}, "Validated")',
        f'=COUNTIF({status_col}{first_row}:{status_col}{last_row}, "Escalated")',
        f'=COUNTIF({status_col}{first_row}:{status_col}{last_row}, "False Positive")',
    ]
    for i, f in enumerate(formulas):
        c = ws.cell(row=5, column=2 + i)
        c.value = f
        c.font = FONT_VALUE
        c.fill = FILL_VALUE
        c.border = BORDER
        c.alignment = AL_CENTER

    # --- header row ---
    for col, text in zip(COLS, HEADERS):
        c = ws[f"{col}8"]
        c.value = text
        c.font = FONT_HEAD
        c.fill = FILL_HEAD
        c.border = BORDER
        c.alignment = AL_LEFT if col in ("D", "H") else AL_CENTER

    # --- data rows ---
    for i, finding in enumerate(findings):
        r = first_row + i
        banded = (r % 2 == 0)
        fill = FILL_BAND if banded else FILL_NONE
        values = {
            "C": finding.get("check_id", ""),
            "D": finding.get("check_title", ""),
            "E": SEVERITY_LABEL.get(finding.get("severity", ""), str(finding.get("severity", "")).capitalize()),
            "F": assign_map.get(finding.get("id"), None),
            "G": status,
            "H": None,
        }
        for col in COLS:
            c = ws[f"{col}{r}"]
            if col in values:
                c.value = values[col]
            if col == "C":
                c.font = FONT_DATA_B
            else:
                c.font = FONT_DATA
            c.border = BORDER
            c.alignment = AL_LEFT if col in ("D", "H") else AL_CENTER
            if col != "B":
                c.fill = fill
        ws.row_dimensions[r].height = 30

    # --- auto numbering in column B (array formula) ---
    if findings:
        ws["B9"] = ArrayFormula(
            f"B{first_row}:B{last_row}",
            f'=IF(C{first_row}:C{last_row}<>"", '
            f'ROW(C{first_row}:C{last_row})-ROW(C{first_row})+1, "")',
        )

    # --- dropdowns ---
    def add_list(formula, ref):
        dv = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(ref)

    add_list(",".join(SEVERITIES), f"E{first_row}:E{last_row}")
    if team:
        add_list(",".join(team), f"F{first_row}:F{last_row}")
    add_list(",".join(STATUSES), f"G{first_row}:G{last_row}")

    # --- layout ---
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[4].height = 19.5
    ws.row_dimensions[5].height = 27
    ws.row_dimensions[8].height = 24
    ws.freeze_panes = "A9"

    return wb, ws


# ---------------------------------------------------------------------------
# CLI helpers
# ---------------------------------------------------------------------------

def resolve_team(args):
    """Combine --assignee (repeatable) and --assignees (comma list)."""
    team = []
    raw = list(args.assignee or [])
    if args.assignees:
        raw.append(args.assignees)
    for item in raw:
        for name in item.split(","):
            name = name.strip()
            if name and name not in team:
                team.append(name)
    if not team:
        team = list(DEFAULT_TEAM)
    if args.assign_all and args.assign_all not in team:
        team.append(args.assign_all)
    return team


def build_assign_map(findings, team, args):
    assign_map = {}
    if args.assign_all:
        for f in findings:
            assign_map[f["id"]] = args.assign_all
    elif args.round_robin and team:
        for i, f in enumerate(findings):
            assign_map[f["id"]] = team[i % len(team)]
    return assign_map


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Build a team assignment Excel tracker from Prowler findings.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  python prowler_excel_tracker.py\n"
            "  python prowler_excel_tracker.py output/prowler-output-XXXX.csv -o tracker.xlsx\n"
            "  python prowler_excel_tracker.py output --assignee KT --assignee Aqmal --round-robin\n"
            "  python prowler_excel_tracker.py output --assignees \"KT,Aqmal,Shalitha\" --assign-all Aqmal\n"
        ),
    )
    parser.add_argument("input", nargs="?", default="output",
                        help="Prowler output file (.html/.csv/.ocsf.json) or a folder (default: ./output)")
    parser.add_argument("-o", "--output", default=None,
                        help="Output .xlsx path (default: <input-dir>/prowler-issues-tracker.xlsx)")
    parser.add_argument("--status", default="FAIL",
                        help="Comma separated statuses to include, or 'all' (default: FAIL)")
    parser.add_argument("--assignee", action="append", metavar="NAME",
                        help="Add a team member to the Assign dropdown (repeatable)")
    parser.add_argument("--assignees", default="",
                        help="Comma separated team members to add to the Assign dropdown")
    parser.add_argument("--assign-all", default="",
                        help="Pre-fill the Assign column with this person for every finding")
    parser.add_argument("--round-robin", action="store_true",
                        help="Distribute findings across the team round-robin")
    parser.add_argument("--title", default=TITLE_TEXT, help="Report title shown in the sheet")
    parser.add_argument("--sheet-name", default="Sheet1", help="Worksheet name")
    parser.add_argument("--open", action="store_true", help="Open the workbook when done")
    args = parser.parse_args(argv)

    # --- load + group findings (reuse prowler_report) ---
    source = pr.resolve_input(args.input)
    records, source_format = pr.load_records(source)

    wanted = pr.parse_statuses(args.status)
    if wanted is not None:
        records = [r for r in records if r["status"] in wanted]
    if not records:
        raise SystemExit(f"[!] No records with status {args.status!r} found in {source.name}")

    findings = pr.group_records(records)

    # --- team + assignment ---
    team = resolve_team(args)
    assign_map = build_assign_map(findings, team, args)

    # --- build + save ---
    wb, ws = build_workbook(
        findings, team, assign_map,
        status=DEFAULT_STATUS, sheet_name=args.sheet_name, title=args.title,
    )

    if args.output:
        out_path = Path(args.output).expanduser()
    else:
        out_path = source.parent / "prowler-issues-tracker.xlsx"
    wb.save(out_path)

    from collections import Counter
    sev = Counter(SEVERITY_LABEL.get(f["severity"], f["severity"]) for f in findings)

    print("=" * 68)
    print("  Prowler issues tracker (Excel)")
    print("=" * 68)
    print(f"  Source          : {source}")
    print(f"  Format          : {source_format}")
    print(f"  Statuses        : {args.status}")
    print(f"  Findings added  : {len(findings)}")
    print(f"  Severity        : " + ", ".join(f"{k}={v}" for k, v in sev.most_common()))
    print(f"  Team (dropdown) : " + ", ".join(team))
    if args.assign_all:
        print(f"  Assigned all to : {args.assign_all}")
    elif args.round_robin:
        print("  Assignment      : round-robin")
    else:
        print("  Assignment      : blank (assign in Excel)")
    print(f"  Output          : {out_path}")
    print("=" * 68)

    if args.open:
        webbrowser.open(out_path.resolve().as_uri())
    return 0


if __name__ == "__main__":
    sys.exit(main())
